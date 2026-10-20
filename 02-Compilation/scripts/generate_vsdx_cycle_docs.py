"""Generate factual release notes in DOCX and PDF. Author: Abdelhalim Talbi."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parents[2]
DOC = ROOT / "05-Documentation"
IMAGE = DOC / "V42.6-VSDX-ARCHITECTURE.png"
BASE = "OpenVisio-V42.6-VSDX-Interop-Guide"
TITLE = "OpenVisio V42.6 — Interopérabilité VSDX (candidat)"
sections = [
("Périmètre et résultats", [
"Correctif de perte de données : conservation des champs metadata, layers et extensions dans la validation, la sauvegarde et l'historique undo/redo.",
"VSDX expérimental : génération d'une archive OPC XML à une page, rectangles et liaisons droites, import d'un sous-ensemble externe, partie JSON OpenVisio embarquée avec contrôle SHA-256.",
"L'interface Tkinter propose importer et exporter VSDX, sans modifier le format natif .ovjson.",
"134 tests Python unitaires exécutés avec succès sur Linux. Aucun test Windows natif ni Microsoft Visio n'a été réalisé." ]),
("Architecture", [
"Le modèle documentaire valide les objets, les liaisons et les calques ; le moteur VSDX adapte la géométrie en pouces et construit les parties OPC et leurs relations.",
"L'import OpenVisio utilise une partie JSON vérifiée par empreinte SHA-256 (sans signature cryptographique) pour contrôler les données du round-trip interne. Un import Visio externe passe par une extraction géométrique limitée." ]),
("Matrice de qualification", [
"Tests unitaires : 134 réussis, 0 échec (Linux / Python).",
"Intégrité du paquet VSDX : ZIP, XML, relations de pages, SHA-256 JSON ; tests locaux réussis.",
"Fidélité interne : égalité stricte du document OpenVisio après export/import ; réussie.",
"Fidélité externe : NON QUALIFIÉE dans Microsoft Visio ; styles, masters, groupes et calques Visio incomplets.",
"Application Windows installable/portable V42.6 : NON COMPILÉE, NON QUALIFIÉE ; les EXE de l'archive restent V42.5.",
"Captures UI réelles : indisponibles (pas de VM Windows graphique)." ]),
("Procédure d'utilisation", [
"Ouvrir OpenVisio ; créer des équipements et des connexions ; sauvegarder en .ovjson.",
"Menu Fichier → Exporter Visio VSDX (expérimental) ; enregistrer le .vsdx.",
"Menu Fichier → Importer Visio VSDX (expérimental) ; vérifier les éléments et enregistrer une copie native .ovjson.",
"Avant échange professionnel, ouvrir et réenregistrer dans Microsoft Visio, vérifier les dimensions, textes et connecteurs et archiver les preuves de fidélité." ]),
("Risques et prochaines actions", [
"Risque P0 : la branche GitHub ne contient pas encore l'intégralité des sources ; la CI Windows n'est pas exécutable en l'état.",
"Risque P1 : conformité OPC/Visio non certifiée ; une ouverture dans Microsoft Visio peut demander une réparation.",
"Action : synchroniser sources + tests sur branche dédiée, faire exécuter CI Windows, puis tests Visio réels et snapshots authentiques.",
"Ne pas fusionner dans main avant qualification. Auteur : Abdelhalim Talbi." ])]

def make_docx():
    d=Document(); sec=d.sections[0]; sec.top_margin=Cm(2); sec.bottom_margin=Cm(1.7)
    d.styles['Normal'].font.name='Calibri'; d.styles['Normal'].font.size=Pt(10)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(TITLE); r.bold=True; r.font.size=Pt(18)
    p=d.add_paragraph('20 octobre 2026 · Auteur : Abdelhalim Talbi · Statut : candidat non qualifié Windows'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for i,(title,bullets) in enumerate(sections):
        d.add_heading(title,level=1)
        if i==1:
            d.add_picture(str(IMAGE),width=Cm(16))
        for bullet in bullets:
            d.add_paragraph(bullet,style='List Bullet')
    d.add_heading('Traçabilité',level=1)
    d.add_paragraph('Tests : 04-Tests/tests/test_layers_persistence_v426.py ; 04-Tests/tests/test_vsdx_v426.py. Sources : 03-Sources/openvisio/validation.py ; 03-Sources/openvisio/vsdx.py ; 03-Sources/openvisio/app.py.')
    d.add_paragraph('Références : Microsoft MS-VSDX — Pages XML Part, Page XML Part (documentation publique).')
    out=DOC/(BASE+'.docx');d.save(out);return out

def make_pdf():
    pdfmetrics.registerFont(TTFont('DejaVu','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
    pdfmetrics.registerFont(TTFont('DejaVuBold','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
    styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='OVTitle',fontName='DejaVuBold',fontSize=15,leading=21,spaceAfter=10))
    styles.add(ParagraphStyle(name='OVHeading',fontName='DejaVuBold',fontSize=12,leading=17,spaceBefore=13,spaceAfter=7))
    styles.add(ParagraphStyle(name='OVBody',fontName='DejaVu',fontSize=9,leading=14,spaceAfter=6))
    story=[Paragraph(TITLE,styles['OVTitle']),Paragraph('20 octobre 2026 | Auteur : Abdelhalim Talbi | Candidat non qualifié Windows',styles['OVBody'])]
    for i,(title,bullets) in enumerate(sections):
        story.append(Paragraph(title,styles['OVHeading']))
        if i==1:story.append(Image(str(IMAGE),width=490,height=150))
        for bullet in bullets:story.append(Paragraph('• '+bullet.replace('&','&amp;').replace('<','&lt;'),styles['OVBody']))
    story.append(Paragraph('Traçabilité',styles['OVHeading']))
    story.append(Paragraph('Sources : validation.py, vsdx.py, app.py. Tests : test_layers_persistence_v426.py, test_vsdx_v426.py.',styles['OVBody']))
    out=DOC/(BASE+'.pdf'); SimpleDocTemplate(str(out),pagesize=A4,rightMargin=44,leftMargin=44,topMargin=45,bottomMargin=45).build(story)
    return out

if __name__=='__main__':print(make_docx(),make_pdf())
