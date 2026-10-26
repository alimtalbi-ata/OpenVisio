#!/usr/bin/env python3
"""Build a small, reviewable DOCX/PDF release brief and render the PDF for QA.
Author: Abdelhalim Talbi. This report distinguishes synthetic tests from real
Windows build evidence and Microsoft Visio Desktop certification.
"""
from __future__ import annotations

import argparse
import json
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer, Table, TableStyle,
    PageBreak, KeepTogether, Flowable,
)
import fitz
from release_document_guard_v4327 import audit_document_artifacts


class PipelineDiagram(Flowable):
    """Compact vector architecture sketch; it is not an application screenshot."""
    def __init__(self, width=170*mm, height=22*mm):
        super().__init__()
        self.width = width
        self.height = height

    def draw(self):
        c = self.canv
        labels = ["ZIP / VSDX", "Private-part guard", "Validated model", "Editor + exports"]
        gap = 4*mm
        box_w = (self.width - gap*3) / 4
        box_h = 13*mm
        y = 4*mm
        for i, label in enumerate(labels):
            x = i * (box_w + gap)
            c.setStrokeColor(colors.HexColor("#375a7f"))
            c.setFillColor(colors.HexColor("#edf4fb"))
            c.roundRect(x, y, box_w, box_h, 2*mm, stroke=1, fill=1)
            c.setFillColor(colors.HexColor("#18324b"))
            c.setFont("Helvetica-Bold", 7.4)
            words = label.split()
            if len(label) > 17:
                mid = max(1, len(words)//2)
                lines = (" ".join(words[:mid]), " ".join(words[mid:]))
            else:
                lines = (label,)
            for n, line in enumerate(lines):
                c.drawCentredString(x+box_w/2, y+box_h/2+(len(lines)-1)*3.5-n*7, line)
            if i < len(labels)-1:
                c.setStrokeColor(colors.HexColor("#375a7f"))
                c.setLineWidth(1.1)
                c.line(x+box_w+0.5*mm, y+box_h/2, x+box_w+gap-0.8*mm, y+box_h/2)
                c.line(x+box_w+gap-2.0*mm, y+box_h/2+1.2*mm, x+box_w+gap-0.8*mm, y+box_h/2)
                c.line(x+box_w+gap-2.0*mm, y+box_h/2-1.2*mm, x+box_w+gap-0.8*mm, y+box_h/2)


def _styles():
    base = getSampleStyleSheet()
    base.add(ParagraphStyle(
        name="OVTitle", parent=base["Title"], fontName="Helvetica-Bold",
        fontSize=19, leading=22, textColor=colors.HexColor("#15324b"),
        alignment=TA_CENTER, spaceAfter=5*mm,
    ))
    base.add(ParagraphStyle(
        name="OVH1", parent=base["Heading1"], fontName="Helvetica-Bold",
        fontSize=12, leading=15, textColor=colors.HexColor("#24537b"),
        spaceBefore=3*mm, spaceAfter=1.8*mm,
    ))
    base.add(ParagraphStyle(
        name="OVBody", parent=base["BodyText"], fontName="Helvetica",
        fontSize=8.6, leading=11.4, spaceAfter=1.6*mm,
    ))
    base.add(ParagraphStyle(
        name="OVSmall", parent=base["BodyText"], fontName="Helvetica",
        fontSize=7.4, leading=9.3, textColor=colors.HexColor("#475569"),
    ))
    return base


def _footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#cbd5e1"))
    canvas.line(18*mm, 14*mm, 192*mm, 14*mm)
    canvas.setFont("Helvetica", 7.3)
    canvas.setFillColor(colors.HexColor("#475569"))
    canvas.drawString(18*mm, 9*mm, "OpenVisio V43.27 candidate | Abdelhalim Talbi")
    canvas.drawRightString(192*mm, 9*mm, f"Page {doc.page}")
    canvas.restoreState()


def _pdf(root: Path, output: Path, report: dict, commit: str):
    styles = _styles()
    output.parent.mkdir(parents=True, exist_ok=True)
    doc = BaseDocTemplate(
        str(output), pagesize=A4, leftMargin=18*mm, rightMargin=18*mm,
        topMargin=15*mm, bottomMargin=19*mm, title="OpenVisio V43.27 Qualification",
        author="Abdelhalim Talbi",
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="body")
    doc.addPageTemplates([PageTemplate(id="all", frames=[frame], onPage=_footer)])
    story = [
        Paragraph("OpenVisio V43.27", styles["OVTitle"]),
        Paragraph("Bounded ShapeSheet scalar formulas + reproducible Windows candidate", styles["OVH1"]),
        Paragraph(
            "Release candidate brief. This package contains the real Python engine source, "
            "tests, build scripts, and Windows executables produced by the associated CI job. "
            "It is not a claim of Microsoft Visio Desktop compatibility.", styles["OVBody"]),
        PipelineDiagram(),
        Spacer(1, 2*mm),
        Paragraph("Functional change", styles["OVH1"]),
        Paragraph(
            "Static route validation keeps its per-segment obstacle prefilter and exact intersection checks. External ShapeSheet geometry now evaluates a bounded scalar subset: arithmetic, numeric cached cell references, numeric ThePage dimensions and an explicit function allowlist. Unsupported, cyclic, non-real, non-finite, over-budget or geometrically invalid results keep cached V values and emit fidelity warnings; strict mode rejects unresolved geometry formulas. Evaluation results are validated before replacing cached geometry, so a formula such as Width=0 cannot silently cause a shape to disappear. No Python eval is used, and this is not a complete ShapeSheet interpreter.", styles["OVBody"]),
        Paragraph(
            "Both mirrored VSDX importers also reject case variants of the reserved "
            "openvisio/document.json and openvisio/manifest.json members. These aliases previously "
            "could bypass embedded-model import and fall back to lossy public geometry, dropping "
            "editor metadata and layers. Path, Win32 alias, collision and decompression-amplification "
            "guards remain active.", styles["OVBody"]),
        Paragraph(
            "The cumulative ZIP validator checks file and directory entries for unsafe paths, "
            "Windows case-fold collisions and file/directory aliases before checking source/test "
            "mirror parity, engine dependencies and complete SHA-256 coverage. CI also extracts "
            "the built ZIP to a clean directory and reruns the full bundled test suite.", styles["OVBody"]),
        Paragraph("Source and dependency layout", styles["OVH1"]),
        Paragraph(
            "03-Sources/openvisio and python/openvisio are compared byte-for-byte (line endings "
            "normalized by the parity gate). The required engine modules include model.py, "
            "vsdx.py, render.py, routing.py, parallel_geometry.py and visio_formula.py. Test mirrors are "
            "04-Tests/tests and python/tests.", styles["OVBody"]),
        Paragraph("Qualification summary", styles["OVH1"]),
    ]
    rows = [["Gate", "Observed result", "Boundary"]]
    unit = report.get("unit_suite", {})
    rows.append(["Python unit/regression suite",
                 f'{unit.get("status", "NOT_RECORDED")} - {unit.get("executed_tests", "unknown")} tests, '
                 f'{unit.get("skipped_tests", "unknown")} skipped',
                 "Synthetic/unit and UI regression fixtures"])
    rows.extend([
        ["Checkout + source parity", report.get("checkout_and_parity", "NOT_RECORDED"),
         "Independent Git checkout; exact source SHA"],
        ["Windows portable build + smoke", report.get("portable_smoke", "NOT_RECORDED"),
         "Frozen application smoke, not manual acceptance"],
        ["Windows installer + smoke", report.get("installer_smoke", "NOT_RECORDED"),
         "Silent installation / installed EXE smoke"],
        ["Microsoft Visio Desktop round-trip", "NOT TESTED",
         "No real Visio Desktop interoperability result"],
        ["Interactive Windows screenshots", "NOT AVAILABLE",
         "Automated smoke only; no manual UI session"],
    ])
    table = Table(rows, colWidths=[42*mm, 47*mm, 83*mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#24537b")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTNAME", (0,1), (-1,-1), "Helvetica"),
        ("FONTSIZE", (0,0), (-1,-1), 7.1),
        ("LEADING", (0,0), (-1,-1), 8.5),
        ("GRID", (0,0), (-1,-1), 0.35, colors.HexColor("#cbd5e1")),
        ("BACKGROUND", (0,1), (-1,-1), colors.white),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING", (0,0), (-1,-1), 4),
        ("RIGHTPADDING", (0,0), (-1,-1), 4),
        ("TOPPADDING", (0,0), (-1,-1), 4),
        ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]))
    story.extend([
        table,
        Spacer(1, 3*mm),
        Paragraph("Candidate provenance", styles["OVH1"]),
        Paragraph(f"Branch: develop/v43.27-formula-driven-vsdx-geometry-20261026", styles["OVBody"]),
        Paragraph(f"Source commit: {commit}", styles["OVBody"]),
        Paragraph(f"Generated on: {date.today().isoformat()}", styles["OVBody"]),
        Paragraph("Known limits", styles["OVH1"]),
        Paragraph(
            "The historical local V43.4 ZIP is not re-labelled as current source: its embedded "
            "engine and dependency tree differ from the GitHub candidate. The cumulative archive "
            "is built from the exact checkout identified above. Internal application version "
            "remains 42.50.0 due to legacy build contracts. No production data is read or changed.",
            styles["OVBody"]),
        Paragraph(
            "Evidence note: the CI run ID, test logs, binary SHA-256 and final archive SHA-256 "
            "are recorded separately in the package evidence and manifest.", styles["OVSmall"]),
    ])
    doc.build(story)


def _docx(output: Path, report: dict, commit: str):
    document = Document()
    document.core_properties.title = "OpenVisio V43.27 - Architecture and Qualification"
    document.core_properties.author = "Abdelhalim Talbi"
    document.core_properties.subject = "VSDX embedded-part case guard and Windows candidate"
    section = document.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(4)
    title = document.add_heading("OpenVisio V43.27", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Bounded ShapeSheet scalar formulas + reproducible Windows candidate")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    p = document.add_paragraph(
        "Release candidate brief. The package includes actual Python engine source, tests, "
        "build scripts and Windows executables produced by CI. It does not claim Microsoft "
        "Visio Desktop compatibility."
    )
    document.add_heading("Functional improvement", level=1)
    document.add_paragraph(
        "Static route validation uses per-segment obstacle prefiltering with exact intersection checks. External ShapeSheet geometry evaluates a bounded scalar subset: arithmetic, numeric cached cell references, numeric ThePage dimensions and an explicit function allowlist. Unsupported, cyclic, non-real, non-finite, over-budget or geometrically invalid formula results keep cached V values and emit fidelity warnings; strict mode rejects unresolved formulas. Result geometry is validated before replacing cached cells. No Python eval is used. The importer also rejects case variants of reserved openvisio/document.json and openvisio/manifest.json members. CI checks source/test parity and reruns the full suite from the extracted ZIP."
    )
    document.add_heading("Source layout", level=1)
    document.add_paragraph(
        "Engine source: 03-Sources/openvisio and python/openvisio. Required modules include "
        "model.py, vsdx.py, render.py, routing.py, parallel_geometry.py and visio_formula.py. Test mirrors: "
        "04-Tests/tests and python/tests. The independent checkout verifier pins the candidate "
        "commit and critical Git blob identifiers."
    )
    document.add_heading("Qualification summary", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Gate", "Observed result", "Boundary"
    unit = report.get("unit_suite", {})
    entries = [
        ("Python unit/regression suite",
         f'{unit.get("status", "NOT_RECORDED")}: {unit.get("executed_tests", "unknown")} tests; '
         f'{unit.get("skipped_tests", "unknown")} skipped',
         "Synthetic/unit and UI regression fixtures"),
        ("Independent checkout and parity", report.get("checkout_and_parity", "NOT_RECORDED"),
         "Exact source SHA and mirrored code"),
        ("Windows portable build/smoke", report.get("portable_smoke", "NOT_RECORDED"),
         "Automated smoke, not manual acceptance"),
        ("Windows installer/smoke", report.get("installer_smoke", "NOT_RECORDED"),
         "Install and launch smoke"),
        ("Microsoft Visio Desktop", "NOT TESTED", "No real round-trip certification"),
        ("Interactive Windows screenshots", "NOT AVAILABLE", "No manual session performed"),
    ]
    for values in entries:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    document.add_heading("Provenance", level=1)
    document.add_paragraph("Branch: develop/v43.27-formula-driven-vsdx-geometry-20261026")
    document.add_paragraph(f"Source commit: {commit}")
    document.add_paragraph(f"Generated on: {date.today().isoformat()}")
    document.add_paragraph("Internal application version remains 42.50.0 due to legacy build contracts.")
    document.add_paragraph(
        "The historical V43.4 ZIP is not relabelled as current source: its engine and dependency "
        "tree differ from the GitHub candidate. The cumulative archive is built from the exact "
        "checkout. No production data is read or changed."
    )
    footer = section.footer.paragraphs[0]
    footer.text = "OpenVisio V43.27 candidate | Author: Abdelhalim Talbi"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--out-dir", type=Path, required=True)
    parser.add_argument("--commit", required=True)
    parser.add_argument("--qualification-json", type=Path, required=True)
    args = parser.parse_args()
    root = args.root.resolve()
    out = args.out_dir.resolve()
    report = json.loads(args.qualification_json.read_text(encoding="utf-8"))
    if report.get("source_commit") != args.commit:
        raise SystemExit("Qualification report commit does not match checked-out source")
    docx_path = out / "OpenVisio-V43.27-Architecture-and-Qualification.docx"
    pdf_path = out / "OpenVisio-V43.27-Architecture-and-Qualification.pdf"
    _docx(docx_path, report, args.commit)
    _pdf(root, pdf_path, report, args.commit)

    # Render every PDF page to PNG and perform layout sanity checks. These
    # previews are CI QA evidence, not Windows application screenshots.
    preview_dir = out / "pdf-page-renders"
    preview_dir.mkdir(parents=True, exist_ok=True)
    pdf = fitz.open(pdf_path)
    if not pdf.page_count:
        raise SystemExit("PDF has no pages")
    visible_chars = 0
    page_renders = []
    for index, page in enumerate(pdf):
        text = page.get_text()
        if len(text.strip()) < 120:
            raise SystemExit(f"PDF page {index+1} contains too little text")
        visible_chars += len(text)
        rect = page.rect
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    x0, y0, x1, y1 = span["bbox"]
                    if x0 < 0 or y0 < 0 or x1 > rect.width + 0.5 or y1 > rect.height + 0.5:
                        raise SystemExit(f"PDF text outside page bounds on page {index+1}")
        pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
        png = preview_dir / f"page-{index+1}.png"
        pix.save(png)
        page_renders.append({"page": index+1, "png": png.name, "visible_characters": len(text)})
    if visible_chars < 500:
        raise SystemExit("PDF text inspection failed")
    qa = {
        "pdf_render_status": "PASS",
        "pdf_pages": pdf.page_count,
        "pdf_visible_characters": visible_chars,
        "pdf_page_renders": page_renders,
        "docx_generated": docx_path.name,
        "docx_structural_status": "PENDING_VALIDATION",
        "docx_visual_render_status": "NOT_PERFORMED_BY_THIS_SCRIPT",
        "note": "PDF pages rendered and checked for text bounds. PNGs are document QA previews, not app screenshots."
    }
    validation = audit_document_artifacts(
        docx_path,
        pdf_path,
        pdf_pages=pdf.page_count,
        pdf_visible_characters=visible_chars,
    )
    qa["docx_structural_status"] = validation["docx_status"]
    qa["pdf_artifact_status"] = validation["pdf_status"]
    qa["docx_bytes"] = validation["docx_bytes"]
    qa["pdf_bytes"] = validation["pdf_bytes"]
    qa["docx_text_characters"] = validation["docx_text_characters"]
    qa["artifact_validation"] = validation
    if validation["pdf_status"] != "PASS":
        qa["pdf_render_status"] = "FAIL"
    (out / "document-qa.json").write_text(json.dumps(qa, indent=2) + "\n", encoding="utf-8")
    if not validation["ok"]:
        raise SystemExit(
            "Generated release document structural QA failed: " + "; ".join(validation["errors"])
        )
    print(json.dumps({"docx": str(docx_path), "pdf": str(pdf_path), "qa": qa}, indent=2))


if __name__ == "__main__":
    main()
